import io
from pathlib import Path
import unittest
from unittest.mock import patch

import cv2
import numpy as np
from openpyxl import load_workbook

import server

FIXTURES = Path(__file__).parent / 'fixtures'


def line(text, x, y, width=.30, height=.04, confidence=95):
    return {'text': text, 'x': x, 'y': y, 'w': width, 'h': height, 'conf': confidence}


class ExtractionRulesTests(unittest.TestCase):
    def extract(self, lines):
        return server.extract_fields(np.zeros((100, 180, 3), dtype=np.uint8),
                                     lines=lines, retry=False)

    def test_split_logo_is_combined_and_checked_against_printed_domain(self):
        fields, _, _ = self.extract([
            line('Ambrish Deshpande', .55, .12), line('COO', .72, .22, .10),
            line('meta', .05, .34, .30, .13), line('I N F O T E C H', .05, .51),
            line('www.metainfotech.com', .53, .83),
        ])
        self.assertEqual(fields['companyName'], 'Meta Infotech')
        self.assertEqual(fields['jobTitle'], 'COO')

    def test_title_zero_correction_requires_the_name_column(self):
        for x, expected in [(.68, 'COO'), (.05, '')]:
            with self.subTest(x=x):
                fields, confidence, _ = self.extract([
                    line('Ambrish Deshpande', .55, .12), line('C00', x, .22, .10),
                ])
                self.assertEqual(fields['jobTitle'], expected)
                if expected:
                    self.assertEqual(confidence['jobTitle'], 'review')

    def test_company_inferred_from_partial_logo_needs_review(self):
        fields, confidence, _ = self.extract([
            line('Jane Doe', .55, .12), line('INFOTECH', .05, .51),
            line('www.metainfotech.com', .53, .83),
        ])
        self.assertEqual(fields['companyName'], 'Meta Infotech')
        self.assertEqual(confidence['companyName'], 'review')

    def test_personal_email_provider_is_not_used_as_the_company(self):
        fields, _, _ = self.extract([
            line('Jane Doe', .55, .12), line('jane@gmail.com', .05, .70),
        ])
        self.assertEqual(fields['companyName'], '')

    def test_clear_company_name_survives_a_different_email_domain(self):
        fields, _, _ = self.extract([
            line('Jane Doe', .55, .12), line('Acme Systems', .05, .15),
            line('jane@example.com', .05, .70), line('www.example.com', .05, .80),
        ])
        self.assertEqual(fields['companyName'], 'Acme Systems')

    def test_alternate_address_reads_do_not_duplicate_lines(self):
        fields, _, _ = self.extract([
            line('Jane Doe', .55, .12),
            line('118 First Floor, Ackruti Star', .05, .60, confidence=95),
            line('118 First Floor, Ackcuti Star', .05, .62, confidence=55),
            line('Andheri East, Mumbai 400093', .05, .70),
            line('Bangalore Hyderabad Delhi Chennai', .05, .80),
            line('Australia Dubai', .05, .90),
        ])
        self.assertEqual(fields['address'], '118 First Floor, Ackruti Star, Andheri East, Mumbai 400093')


@unittest.skipUnless(all((FIXTURES / name).exists() for name in
                        ['gold-card.jpg', 'white-card.jpg', 'three-cards.jpg']),
                     'Local card-photo fixtures are excluded from GitHub; copy them into tests/fixtures to run photo checks.')
class PhotoRegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cards = {name: server.process_image((FIXTURES / name).read_bytes())
                     for name in ['gold-card.jpg', 'white-card.jpg', 'three-cards.jpg']}

    def test_single_textured_card_keeps_its_name_title_and_company(self):
        cards = self.cards['gold-card.jpg']
        self.assertEqual(len(cards), 1)
        self.assertNotIn('error', cards[0])
        fields = cards[0]['fields']
        self.assertEqual(fields['name'], 'Ambrish Deshpande')
        self.assertEqual(fields['jobTitle'], 'COO')
        self.assertEqual(fields['companyName'], 'Meta Infotech')
        self.assertEqual(fields['email'], 'ambrishd@metainfotech.com')

    def test_white_card_still_extracts_correctly(self):
        cards = self.cards['white-card.jpg']
        self.assertEqual(len(cards), 1)
        fields = cards[0]['fields']
        self.assertEqual(fields['name'], 'Pavan Kumar Verma')
        self.assertEqual(fields['companyName'], 'Embee')
        self.assertEqual(fields['jobTitle'], 'Director Solution Sales - MWS')

    def test_three_cards_are_separate_and_ordered_with_independent_fields(self):
        cards = self.cards['three-cards.jpg']
        self.assertEqual(len(cards), 3)
        expected = [
            ('Ambrish Deshpande', 'Meta Infotech', 'COO', 'ambrishd@metainfotech.com'),
            ('Dr. Nishant Sawant', 'Meta Infotech', 'Director - Managed Security Services', 'nishant@metainfotech.com'),
            ('Pavan Kumar Verma', 'Embee', 'Director Solution Sales - MWS', 'pavan.verma@embee.co.in'),
        ]
        for index, (card, values) in enumerate(zip(cards, expected), 1):
            self.assertNotIn('error', card)
            self.assertEqual(card['index'], index)
            self.assertEqual(tuple(card['fields'][key] for key in
                                   ['name', 'companyName', 'jobTitle', 'email']), values)
        self.assertEqual(len({card['preview'] for card in cards}), 3)

    def test_all_three_cards_export_as_independent_excel_rows(self):
        rows = [{**card['fields'], 'sourceImage': 'three-cards.jpg'}
                for card in self.cards['three-cards.jpg']]
        workbook = load_workbook(io.BytesIO(server.export_xlsx(rows)))
        sheet = workbook.active
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual([sheet[f'B{i}'].value for i in [2, 3, 4]],
                         ['Ambrish Deshpande', 'Dr. Nishant Sawant', 'Pavan Kumar Verma'])
        self.assertEqual([sheet[f'D{i}'].value for i in [2, 3, 4]],
                         ['COO', 'Director - Managed Security Services', 'Director Solution Sales - MWS'])
        self.assertEqual([sheet[f'J{i}'].value for i in [2, 3, 4]], ['three-cards.jpg'] * 3)

    def test_text_fallback_continues_when_only_one_rectangle_is_found(self):
        image = cv2.imread(str(FIXTURES / 'three-cards.jpg'))
        quads = server.detect_quads(image)
        with patch.object(server, 'detect_quads', return_value=[quads[-1]]):
            crops = server.detect_cards(image)
        self.assertEqual(len(crops), 3)

    def test_cards_in_a_row_or_grid_are_not_merged_or_duplicated(self):
        source = cv2.imread(str(FIXTURES / 'three-cards.jpg'))
        crops = server.detect_cards(source)
        thumbnails = [cv2.resize(crop, (520, 310)) for crop in crops]
        for positions, shape in [
            ([(30, 30), (610, 30), (1190, 30)], (380, 1740)),
            ([(30, 30), (610, 30), (30, 400), (610, 400)], (750, 1160)),
        ]:
            with self.subTest(cards=len(positions)):
                board = np.full((*shape, 3), 110, dtype=np.uint8)
                for index, (x, y) in enumerate(positions):
                    board[y:y+310, x:x+520] = thumbnails[index % len(thumbnails)]
                self.assertEqual(len(server.detect_cards(board)), len(positions))


if __name__ == '__main__':
    unittest.main()
