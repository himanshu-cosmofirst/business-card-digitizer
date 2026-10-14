import base64
import io
from pathlib import Path
import threading
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from app import create_app


class HostedAppTests(unittest.TestCase):
    def setUp(self):
        self.app = create_app({'TESTING': True, 'APP_PASSWORD': ''})
        self.client = self.app.test_client()

    def test_page_and_health_are_served_by_the_hosted_app(self):
        page = self.client.get('/')
        self.assertEqual(page.status_code, 200)
        self.assertIn(b'Business Card Digitizer', page.data)
        page.close()
        response = self.client.get('/api/health')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['app'], 'business-card-digitizer-v2')
        self.assertIn('Tesseract', response.json['ocr'])

    def test_route_method_and_bad_json_errors_are_json(self):
        for method, path, kwargs, status in [
            ('GET', '/api/missing', {}, 404),
            ('GET', '/api/process', {}, 405),
            ('POST', '/api/process', {'data': '<html>', 'content_type': 'application/json'}, 400),
            ('POST', '/api/process', {'json': []}, 400),
            ('POST', '/api/process', {'json': {'data': '%%%'}}, 400),
            ('POST', '/api/export', {'json': {'rows': []}}, 400),
            ('POST', '/api/export', {'json': {'rows': 'not-cards'}}, 400),
        ]:
            with self.subTest(path=path, status=status):
                response = self.client.open(path, method=method, **kwargs)
                self.assertEqual(response.status_code, status)
                self.assertTrue(response.is_json)
                self.assertFalse(response.json['ok'])

    def test_oversize_request_has_a_json_explanation(self):
        self.app.config['MAX_CONTENT_LENGTH'] = 16
        response = self.client.post('/api/process', json={'data': 'a' * 100})
        self.assertEqual(response.status_code, 413)
        self.assertIn('too large', response.json['error'])

    def test_password_protects_the_page_and_uploads_but_allows_health_checks(self):
        app = create_app({'TESTING': True, 'APP_USERNAME': 'admin', 'APP_PASSWORD': 'test-password'})
        client = app.test_client()
        for path in ['/', '/api/process', '/api/export']:
            response = client.get(path) if path == '/' else client.post(path, json={})
            self.assertEqual(response.status_code, 401)
            self.assertIn('Basic', response.headers['WWW-Authenticate'])
        self.assertEqual(client.get('/api/health').status_code, 200)
        auth = 'Basic ' + base64.b64encode(b'admin:test-password').decode()
        page = client.get('/', headers={'Authorization': auth})
        self.assertEqual(page.status_code, 200)
        page.close()
        invalid = 'Basic ' + base64.b64encode(b'admin:wrong-password').decode()
        self.assertEqual(client.get('/', headers={'Authorization': invalid}).status_code, 401)

    def test_busy_ocr_returns_json_and_releases_its_slot(self):
        started, finish = threading.Event(), threading.Event()
        responses = []

        def slow_process(_):
            started.set()
            self.assertTrue(finish.wait(5))
            return []

        def upload():
            with self.app.test_client() as client:
                responses.append(client.post('/api/process', json={'data': 'aW1hZ2U='}))

        with patch('app.process_image', side_effect=slow_process):
            thread = threading.Thread(target=upload)
            thread.start()
            try:
                self.assertTrue(started.wait(5))
                busy = self.client.post('/api/process', json={'data': 'aW1hZ2U='})
                self.assertEqual(busy.status_code, 503)
                self.assertIn('Another photo', busy.json['error'])
            finally:
                finish.set()
                thread.join(5)
        self.assertEqual(responses[0].status_code, 200)
        with patch('app.process_image', return_value=[]):
            self.assertEqual(self.client.post('/api/process', json={'data': 'aW1hZ2U='}).status_code, 200)

    @unittest.skipUnless((Path(__file__).parent / 'fixtures' / 'three-cards.jpg').exists(),
                         'Local photo fixtures are excluded from GitHub.')
    def test_real_three_card_upload_and_export_work_through_hosted_routes(self):
        photo = (Path(__file__).parent / 'fixtures' / 'three-cards.jpg').read_bytes()
        response = self.client.post('/api/process', json={
            'name': 'three-cards.jpg',
            'data': 'data:image/jpeg;base64,' + base64.b64encode(photo).decode(),
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['count'], 3)
        names = [card['fields']['name'] for card in response.json['cards']]
        self.assertEqual(names, ['Ambrish Deshpande', 'Dr. Nishant Sawant', 'Pavan Kumar Verma'])
        rows = [{**card['fields'], 'sourceImage': 'three-cards.jpg'}
                for card in response.json['cards']]
        exported = self.client.post('/api/export', json={'rows': rows})
        self.assertEqual(exported.status_code, 200)
        workbook = load_workbook(io.BytesIO(exported.data))
        self.assertEqual(workbook.active.max_row, 4)
        self.assertEqual([workbook.active[f'B{i}'].value for i in [2, 3, 4]], names)
        exported.close()


if __name__ == '__main__':
    unittest.main()
