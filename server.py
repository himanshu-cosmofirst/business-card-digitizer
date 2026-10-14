import base64, io, json, os, re, sys, traceback, webbrowser, threading
from difflib import SequenceMatcher
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from urllib.parse import urlparse

import cv2
import numpy as np
from PIL import Image
import pytesseract
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.abspath(__file__))
# Windows convenience: standard Tesseract installer location, even if PATH was not updated.
if os.name == 'nt':
    common = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
    if os.path.exists(common):
        pytesseract.pytesseract.tesseract_cmd = common
MAX_BODY = 30 * 1024 * 1024

EMAIL_RE = re.compile(r"[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}", re.I)
WEB_RE = re.compile(r"(?:https?://)?(?:www\.)?[A-Z0-9\-]+(?:\.[A-Z0-9\-]+)+(?:(?:/[^\s|]*)?)", re.I)
PHONE_RE = re.compile(r"(?:\+?\d[\d\s().\-]{7,}\d)")
JOB_WORDS = {
    'director','manager','executive','officer','engineer','developer','designer','sales','marketing','head',
    'president','vice','vp','ceo','coo','cfo','cto','founder','partner','consultant','specialist','lead',
    'solution','solutions','business','operations','operation','procurement','purchase','hr','admin','general',
    'regional','national','senior','sr','assistant','associate','chairman','proprietor','owner','architect',
    'cio','cmo','cso','chro','md','security'
}
TITLE_ACRONYMS = {'CEO','COO','CFO','CTO','CIO','CMO','CSO','CHRO','VP','MD','HR'}
PERSONAL_DOMAINS = {'gmail','yahoo','outlook','hotmail','live','icloud','aol','proton','protonmail'}
COMPANY_WORDS = {
    'pvt','private','limited','ltd','llp','inc','corp','corporation','company','co','technologies','technology',
    'infotech','solutions','industries','industry','enterprises','enterprise','group','systems','services','india',
    'labs','consulting','consultants','international','global','digital','software','communications','ventures'
}
ADDRESS_WORDS = {
    'floor','road','rd','street','st','lane','plot','sector','suite','suites','centre','center','district','building',
    'tower','office','near','opposite','opp','mumbai','delhi','gurgaon','gurugram','noida','bangalore','bengaluru',
    'hyderabad','chennai','kolkata','pune','india','maharashtra','nagar','andheri','patna','raipur','jamshedpur',
    'thane','australia','dubai'
}
LABEL_WORDS = {'phone','mobile','mob','email','mail','website','web','tel','telephone','fax','www'}


def clean_space(s):
    return re.sub(r'\s+', ' ', s or '').strip(' \t\r\n|,;:')


def camel_split(s):
    if ' ' not in s and re.search(r'[a-z][A-Z]', s):
        return re.sub(r'(?<=[a-z])(?=[A-Z])', ' ', s)
    return s


def normalize_url(s):
    s = clean_space(s).rstrip('.,;)')
    if re.match(r'^ww\.', s, re.I): s='www.'+s[3:]
    s = re.sub(r'^(?:website|web)\s*[:\-]?\s*', '', s, flags=re.I)
    return s


def normalize_phone(s):
    s = clean_space(s)
    s = re.sub(r'^(?:phone|mobile|mob|tel|telephone)\s*[:\-]?\s*', '', s, flags=re.I)
    # keep +, digits and readable spaces/hyphens
    s = re.sub(r'[^\d+()\-\s]', '', s)
    return clean_space(s)


def order_points(pts):
    pts = np.asarray(pts, dtype=np.float32)
    rect = np.zeros((4,2), dtype=np.float32)
    s = pts.sum(axis=1)
    d = np.diff(pts, axis=1).ravel()
    rect[0] = pts[np.argmin(s)]
    rect[2] = pts[np.argmax(s)]
    rect[1] = pts[np.argmin(d)]
    rect[3] = pts[np.argmax(d)]
    return rect


def warp_quad(img, pts):
    rect = order_points(pts)
    tl,tr,br,bl = rect
    max_w = int(max(np.linalg.norm(br-bl), np.linalg.norm(tr-tl)))
    max_h = int(max(np.linalg.norm(tr-br), np.linalg.norm(tl-bl)))
    if max_w < 50 or max_h < 30:
        return None
    dst = np.array([[0,0],[max_w-1,0],[max_w-1,max_h-1],[0,max_h-1]], dtype=np.float32)
    M = cv2.getPerspectiveTransform(rect, dst)
    warped = cv2.warpPerspective(img, M, (max_w, max_h))
    if warped.shape[0] > warped.shape[1]:
        warped = cv2.rotate(warped, cv2.ROTATE_90_CLOCKWISE)
    return warped


def rect_iou(a,b):
    ax,ay,aw,ah=a; bx,by,bw,bh=b
    x1=max(ax,bx); y1=max(ay,by); x2=min(ax+aw,bx+bw); y2=min(ay+ah,by+bh)
    inter=max(0,x2-x1)*max(0,y2-y1)
    if not inter: return 0.0
    return inter/(aw*ah+bw*bh-inter)


def rect_overlap(a,b):
    """Intersection as a fraction of the smaller region, for duplicate crops."""
    ax,ay,aw,ah=a; bx,by,bw,bh=b
    inter=max(0,min(ax+aw,bx+bw)-max(ax,bx))*max(0,min(ay+ah,by+bh)-max(ay,by))
    return inter/max(1,min(aw*ah,bw*bh))


def detect_quads(img):
    h,w = img.shape[:2]
    scale = min(1.0, 1800.0/max(h,w))
    sm = cv2.resize(img, None, fx=scale, fy=scale, interpolation=cv2.INTER_AREA) if scale < 1 else img.copy()
    sh,sw = sm.shape[:2]
    gray = cv2.cvtColor(sm, cv2.COLOR_BGR2GRAY)
    gray = cv2.GaussianBlur(gray,(5,5),0)
    edges = cv2.Canny(gray, 35, 120)
    edges = cv2.morphologyEx(edges, cv2.MORPH_CLOSE, np.ones((5,5),np.uint8), iterations=2)
    # Coloured cards can have almost no grayscale contrast against the background.
    blur = cv2.GaussianBlur(sm,(9,9),0)
    lab = cv2.cvtColor(blur,cv2.COLOR_BGR2LAB)
    hsv = cv2.cvtColor(blur,cv2.COLOR_BGR2HSV)
    masks = [edges]
    for channel in (lab[:,:,1],lab[:,:,2],hsv[:,:,1]):
        boundary=cv2.Canny(channel,18,55)
        masks.append(cv2.morphologyEx(boundary,cv2.MORPH_CLOSE,np.ones((7,7),np.uint8),iterations=2))
    contours=[]
    for mask in masks:
        found,_=cv2.findContours(mask,cv2.RETR_LIST,cv2.CHAIN_APPROX_SIMPLE)
        contours.extend(found)
    items=[]
    min_area = max(8000, sw*sh*0.012)
    for c in sorted(contours, key=cv2.contourArea, reverse=True):
        area=cv2.contourArea(c)
        if area < min_area or area > sw*sh*0.92: continue
        peri=cv2.arcLength(c,True)
        approx=None
        for eps in (0.015,0.02,0.025,0.03,0.04):
            a=cv2.approxPolyDP(c,eps*peri,True)
            if len(a)==4 and cv2.isContourConvex(a):
                approx=a; break
        if approx is None: continue
        x,y,bw,bh=cv2.boundingRect(approx)
        sides=[np.linalg.norm(approx[i,0]-approx[(i+1)%4,0]) for i in range(4)]
        card_w=(sides[0]+sides[2])/2; card_h=(sides[1]+sides[3])/2
        ar=max(card_w,card_h)/max(1,min(card_w,card_h))
        fill=area/max(1,bw*bh)
        if not (1.25 <= ar <= 2.35 and fill >= 0.60): continue
        r=(int(x/scale),int(y/scale),int(bw/scale),int(bh/scale))
        pts=approx.reshape(4,2)/scale
        # A loose contour including the background can be larger than the real card.
        items.append({'rect':r,'pts':pts,'score':float(area*fill**4)})
    selected=[]
    for item in sorted(items,key=lambda q:q['score'],reverse=True):
        if any(rect_iou(item['rect'],old['rect'])>.60 or rect_overlap(item['rect'],old['rect'])>.85 for old in selected): continue
        selected.append(item)
    return sorted(selected,key=lambda q:(q['rect'][1],q['rect'][0]))


def text_cluster_regions(img, excluded=()):
    """Fallback: use OCR word locations to find separated text islands (useful for many cards in one photo)."""
    h,w=img.shape[:2]
    scale=min(1.0, 2200.0/max(h,w))
    sm=cv2.resize(img,None,fx=scale,fy=scale,interpolation=cv2.INTER_AREA) if scale<1 else img
    try:
        d=pytesseract.image_to_data(sm, config='--oem 3 --psm 11', output_type=pytesseract.Output.DICT)
    except Exception:
        return []
    mask=np.zeros(sm.shape[:2],np.uint8)
    boxes=[]; contacts=[]; web_contacts=[]
    for i,t in enumerate(d['text']):
        t=clean_space(t)
        try: conf=float(d['conf'][i])
        except: conf=-1
        if len(t)<2 or conf<25: continue
        x,y,bw,bh=[int(d[k][i]) for k in ('left','top','width','height')]
        if bw<3 or bh<3: continue
        # Known card rectangles must not bridge text clusters on other cards.
        cx=(x+bw/2)/scale; cy=(y+bh/2)/scale
        if any(rx<=cx<=rx+rw and ry<=cy<=ry+rh for rx,ry,rw,rh in excluded): continue
        boxes.append((x,y,bw,bh))
        if '@' in t: contacts.append(((x+bw/2)/scale,(y+bh/2)/scale))
        elif t.lower().startswith('www.'): web_contacts.append(((x+bw/2)/scale,(y+bh/2)/scale))
        cv2.rectangle(mask,(x,y),(x+bw,y+bh),255,-1)
    if len(boxes)<4: return []
    anchors=contacts if len(contacts)>=2 else web_contacts
    # connect words within the same card, but not across normal gaps between cards
    kx=max(17, int(sm.shape[1]*0.018)); ky=max(15,int(sm.shape[0]*0.025))
    mask=cv2.dilate(mask,np.ones((ky,kx),np.uint8),iterations=1)
    mask=cv2.morphologyEx(mask,cv2.MORPH_CLOSE,np.ones((max(9,ky//2),max(9,kx//2)),np.uint8),iterations=2)
    cs,_=cv2.findContours(mask,cv2.RETR_EXTERNAL,cv2.CHAIN_APPROX_SIMPLE)
    regions=[]
    for c in cs:
        x,y,bw,bh=cv2.boundingRect(c)
        if bw*bh < sm.shape[0]*sm.shape[1]*0.001: continue
        regions.append((int(x/scale),int(y/scale),int(bw/scale),int(bh/scale)))
    # Merge nearby text islands that belong to the same card (e.g. left/right columns).
    def gap(a,b):
        ax,ay,aw,ah=a; bx,by,bw,bh=b
        gx=max(0,max(ax,bx)-min(ax+aw,bx+bw))
        gy=max(0,max(ay,by)-min(ay+ah,by+bh))
        ox=min(ax+aw,bx+bw)-max(ax,bx)
        oy=min(ay+ah,by+bh)-max(ay,by)
        return gx,gy,ox,oy
    changed=True
    while changed and len(regions)>1:
        changed=False
        for i in range(len(regions)):
            if changed: break
            for j in range(i+1,len(regions)):
                a,b=regions[i],regions[j]; gx,gy,ox,oy=gap(a,b)
                # Join columns before padding; padding first joins separate cards.
                close=(gx <= int(w*0.07) and oy>min(a[3],b[3])*.25) or (gy <= min(a[2],b[2])*.25 and ox>min(a[2],b[2])*.35)
                if not close: continue
                x=min(a[0],b[0]); y=min(a[1],b[1]); x2=max(a[0]+a[2],b[0]+b[2]); y2=max(a[1]+a[3],b[1]+b[3])
                u=(x,y,x2-x,y2-y)
                contact_points=[(cx,cy) for cx,cy in anchors if x<=cx<=x2 and y<=cy<=y2]
                if len(contact_points)>=2:
                    xs,ys=zip(*contact_points)
                    # Two distant contact blocks are evidence of separate cards.
                    if max(ys)-min(ys)>u[2]*.22 and u[3]>u[2]*.55: continue
                    if max(xs)-min(xs)>u[3]*.8 and u[2]>u[3]*2.8: continue
                if 1.05 <= u[2]/max(1,u[3]) <= 4.5:
                    regions[i]=u; regions.pop(j); changed=True; break
    results=[]
    for x,y,bw,bh in regions:
        words=[box for box in boxes if x<=box[0]/scale<=x+bw and y<=box[1]/scale<=y+bh]
        if len(words)<6 or not 1.05<=bw/max(1,bh)<=3.0: continue
        # OCR may only recognize the start of a faint name/email. Leave enough
        # horizontal margin to avoid cutting off its unrecognized final letters.
        px=int(bw*.16); py=int(bh*.10)
        x1=max(0,x-px); y1=max(0,y-py)
        results.append((x1,y1,min(w,x+bw+px)-x1,min(h,y+bh+py)-y1))
    return sorted(results,key=lambda r:(r[1],r[0]))


def crop_region(img, rect, pad=.03):
    x,y,w,h=rect; H,W=img.shape[:2]
    px=int(w*pad); py=int(h*pad)
    x=max(0,x-px); y=max(0,y-py); x2=min(W,x+w+2*px); y2=min(H,y+h+2*py)
    crop=img[y:y2,x:x2]
    if crop.shape[0] > crop.shape[1]: crop=cv2.rotate(crop,cv2.ROTATE_90_CLOCKWISE)
    return crop


def detect_cards(img):
    quads=detect_quads(img)
    regions=[]
    for q in quads:
        crop=warp_quad(img,q['pts'])
        if crop is not None: regions.append((q['rect'],crop))
    # Always check outside the detected rectangles, even when one card was found.
    for rect in text_cluster_regions(img,excluded=[rect for rect,_ in regions]):
        if any(rect_overlap(rect,old)>.55 for old,_ in regions): continue
        regions.append((rect,crop_region(img,rect,0)))
    regions.sort(key=lambda item:(item[0][1],item[0][0]))
    cards=[crop for _,crop in regions if crop.size]
    # Single-card fallback: OCR whole photo; this is deliberate so difficult backgrounds still work.
    if not cards:
        cards=[img]
    return cards


def enhance_for_ocr(img):
    # Keep original colour/texture. Aggressive thresholding hurts textured business cards.
    h,w=img.shape[:2]
    scale=min(1.0, 2400.0/max(h,w))
    return cv2.resize(img,None,fx=scale,fy=scale,interpolation=cv2.INTER_AREA) if scale<1 else img.copy()

def ink_image(img):
    """An additional OCR view that suppresses light texture, preserving the original."""
    gray=cv2.cvtColor(img,cv2.COLOR_BGR2GRAY) if img.ndim==3 else img
    gray=cv2.medianBlur(gray,3)
    if np.median(gray)<100: gray=255-gray
    otsu,_=cv2.threshold(gray,0,255,cv2.THRESH_BINARY+cv2.THRESH_OTSU)
    threshold=min(otsu,float(np.median(gray))*.72)
    binary=cv2.threshold(gray,threshold,255,cv2.THRESH_BINARY)[1]
    if img.ndim==3:
        lab=cv2.cvtColor(cv2.medianBlur(img,3),cv2.COLOR_BGR2LAB).astype(np.int16)
        # Faded coloured ink may be as bright as the card, but its chroma differs.
        colour_distance=np.abs(lab[:,:,1]-np.median(lab[:,:,1]))+np.abs(lab[:,:,2]-np.median(lab[:,:,2]))
        coloured=(colour_distance>28)&(gray<float(np.median(gray))*.97)
        binary[coloured]=0
    return binary


def chroma_image(img):
    """Separate faded coloured lettering from warm/cool card stock."""
    if img.ndim!=3: return None
    lab=cv2.cvtColor(img,cv2.COLOR_BGR2LAB)
    medians=[float(np.median(lab[:,:,channel])) for channel in (1,2)]
    index=max(range(2),key=lambda i:abs(medians[i]-128))
    median=medians[index]
    if abs(median-128)<10: return None
    channel=lab[:,:,index+1]
    foreground=channel<median-10 if median>128 else channel>median+10
    return cv2.medianBlur(np.where(foreground,0,255).astype(np.uint8),3)


def tesseract_lines(img, psm=6, cleaned=False, chroma=False):
    prep=enhance_for_ocr(img)
    if cleaned: prep=ink_image(prep)
    if chroma:
        prep=chroma_image(prep)
        if prep is None: return []
    d=pytesseract.image_to_data(prep, config=f'--oem 3 --psm {psm}', output_type=pytesseract.Output.DICT)
    groups={}
    for i,t in enumerate(d['text']):
        t=clean_space(t)
        if not t: continue
        try: conf=float(d['conf'][i])
        except: conf=-1
        if conf < 15: continue
        key=(d['block_num'][i],d['par_num'][i],d['line_num'][i])
        g=groups.setdefault(key, {'words':[],'conf':[],'x':10**9,'y':10**9,'x2':0,'y2':0})
        g['words'].append(t); g['conf'].append(conf)
        x,y,bw,bh=[int(d[k][i]) for k in ('left','top','width','height')]
        g['x']=min(g['x'],x); g['y']=min(g['y'],y); g['x2']=max(g['x2'],x+bw); g['y2']=max(g['y2'],y+bh)
    lines=[]
    H,W=prep.shape[:2]
    for g in groups.values():
        txt=clean_space(' '.join(g['words']))
        if not txt or sum(ch.isalnum() for ch in txt)<2: continue
        lines.append({'text':txt,'conf':sum(g['conf'])/len(g['conf']), 'x':g['x']/W, 'y':g['y']/H, 'w':(g['x2']-g['x'])/W, 'h':(g['y2']-g['y'])/H})
    lines.sort(key=lambda z:(z['y'],z['x']))
    return lines


def merge_ocr_lines(lines):
    merged=[]
    for ln in sorted(lines,key=lambda line:line['conf'],reverse=True):
        key=re.sub(r'\W+','',ln['text']).lower()
        if len(key)<2: continue
        # Preserve different readings; exact duplicates keep the strongest evidence.
        if any(key==re.sub(r'\W+','',old['text']).lower() and abs(ln['y']-old['y'])<.10 for old in merged): continue
        merged.append(ln)
    return sorted(merged,key=lambda line:(line['y'],line['x']))


def title_text(text,near_name=False):
    text=clean_space(text).strip(' _—|,:;')
    compact=re.sub(r'[\s.]','',text).upper()
    if compact in TITLE_ACRONYMS: return compact
    if near_name and re.fullmatch(r'[A-Z0]{2,5}',compact) and '0' in compact:
        repaired=compact.replace('0','O')
        if repaired in TITLE_ACRONYMS: return repaired
    return text


def same_column(a,b):
    overlap=max(0,min(a['x']+a['w'],b['x']+b['w'])-max(a['x'],b['x']))
    return overlap/max(.001,min(a['w'],b['w']))>.35


def company_from_domain(website):
    host=website.lower().split('://')[-1].split('/')[0].removeprefix('www.')
    parts=host.split('.')
    if len(parts)<2: return ''
    stem=parts[-3] if len(parts)>=3 and '.'.join(parts[-2:]) in {'co.in','co.uk','com.au','co.nz','co.za'} else parts[-2]
    if stem in PERSONAL_DOMAINS: return ''
    stem=stem.replace('-',' ')
    for suffix in sorted(COMPANY_WORDS,key=len,reverse=True):
        if len(suffix)>3 and stem.endswith(suffix) and stem!=suffix:
            stem=stem[:-len(suffix)]+' '+suffix; break
    return ' '.join(word.capitalize() for word in stem.split())



def likely_name(line):
    s=camel_split(clean_space(line['text']))
    low=s.lower()
    if re.search(r'[^A-Za-z\s.\-\'’]',s.strip(' _—|,:;')): return -99
    if any(ch.isdigit() for ch in s) or '@' in s or WEB_RE.search(s) or 'www' in low or len(s)>45: return -99
    words=re.findall(r"[A-Za-z][A-Za-z.'-]+",s)
    if not (2<=len(words)<=4): return -99
    if any(w.lower().strip('.') in JOB_WORDS|COMPANY_WORDS|ADDRESS_WORDS|LABEL_WORDS for w in words): return -20
    score=0
    title=sum(1 for w in words if w[:1].isupper())
    if line['conf']<40 or title<2: return -99
    score += title*3
    if line['y']<.45: score+=4
    if .025<line['h']<.14: score+=min(6,line['h']*80)
    score += line['conf']/10
    if all(len(w)>1 for w in words): score+=2
    return score


def extract_fields(img, lines=None, retry=True):
    if lines is None:
        # Sparse text finds independent columns and logos; the cleaned view recovers
        # small titles on textured stock. Block OCR supplies alternate contact reads.
        lines=tesseract_lines(img,11)+tesseract_lines(img,11,cleaned=True)+tesseract_lines(img,6)
    merged=merge_ocr_lines(lines)
    if retry and not any(likely_name(ln)>=20 and ln['y']<.5 for ln in merged):
        # Identity text can be printed in faded coloured ink even when contacts
        # are black. Keep this extra pass out of email/phone recognition.
        identity=[ln for ln in tesseract_lines(img,11,chroma=True)
                  if ln['y']<.5 and '@' not in ln['text'] and not any(ch.isdigit() for ch in ln['text'])]
        merged=merge_ocr_lines(merged+identity)
    raw='\n'.join(l['text'] for l in merged)
    corpus=raw

    emails=[]
    for e in EMAIL_RE.findall(corpus):
        e=e.strip('.,;:()[]').lower()
        if e not in emails: emails.append(e)
    websites=[]
    for match in WEB_RE.finditer(corpus):
        # Do not treat a substring of an email address as a printed website.
        if match.start()>0 and corpus[match.start()-1]=='@': continue
        u=normalize_url(match.group()).lower()
        if '@' in u: continue
        if re.search(r'\.(com|co\.in|in|org|net|io|ai|biz|co|info)$',u,re.I) and u not in websites:
            websites.append(u)
    # Prefer explicit www lines
    websites.sort(key=lambda u:(0 if u.startswith('www.') else 1, len(u)))
    website=websites[0] if websites else ''
    website_observed=bool(website); email_repaired=False
    # Repair slightly misread email domain when website domain is strongly available.
    if emails and website:
        domain=website.lower().split('://')[-1].removeprefix('www.').split('/')[0]
        emails.sort(key=lambda email:(email.split('@')[-1]!=domain,-corpus.lower().count(email)))
        local,_,ed=emails[0].partition('@')
        if ed and ed!=domain:
            # modest edit-distance heuristic
            def lev(a,b):
                dp=list(range(len(b)+1))
                for i,ca in enumerate(a,1):
                    nd=[i]
                    for j,cb in enumerate(b,1): nd.append(min(nd[-1]+1,dp[j]+1,dp[j-1]+(ca!=cb)))
                    dp=nd
                return dp[-1]
            if lev(ed,domain)<=3:
                emails[0]=local+'@'+domain; email_repaired=True
    if not website and emails:
        website='www.'+emails[0].split('@',1)[1]

    phone_candidates=[]
    for ln in merged:
        for p in PHONE_RE.findall(ln['text']):
            digits=re.sub(r'\D','',p)
            if 9<=len(digits)<=14:
                phone_candidates.append((normalize_phone(p),ln['text'].lower()))
    # dedupe by digits
    ded=[]; seenp=set()
    for p,ctx in phone_candidates:
        d=re.sub(r'\D','',p)
        if d not in seenp: seenp.add(d); ded.append((p,ctx))
    personal=''; company_phone=''
    for p,ctx in ded:
        if any(k in ctx for k in ('mobile','mob')) and not personal: personal=p
        elif any(k in ctx for k in ('phone','tel','office')) and not company_phone: company_phone=p
    if not personal and ded: personal=ded[0][0]
    if not company_phone and len(ded)>1:
        company_phone=next((p for p,_ in ded if p!=personal),'')

    # Name scoring
    name=''; best=0; name_line=None
    for ln in merged:
        sc=likely_name(ln)
        if sc>best:
            best=sc; name=camel_split(clean_space(ln['text'])); name_line=ln
    # remove leading OCR punctuation
    name=re.sub(r'^[^A-Za-z]+','',name).strip()

    def choose_job(candidates):
        job=''; score_best=0; evidence=None
        for ln in candidates:
            near=name_line is not None and same_column(name_line,ln) and 0<ln['y']-name_line['y']<.22
            s=title_text(ln['text'],near)
            low=s.lower(); words=set(re.findall(r'[a-z]+',low))
            if '@' in s or any(ch.isdigit() for ch in s) or 'www' in low: continue
            hits=len(words & JOB_WORDS)
            score=hits*7+(7 if near else 0)+ln['conf']/25
            if score>score_best and hits:
                score_best=score; job=s; evidence=ln
        return job,score_best,evidence

    job,job_score,job_line=choose_job(merged)
    if retry and job_line is not None:
        name_candidates=[ln for ln in merged if likely_name(ln)>0 and same_column(ln,job_line)
                         and 0<job_line['y']-ln['y']<.22]
        variants={re.sub(r'\W+','',ln['text']).lower() for ln in name_candidates}
        needs_name_retry=(not name_line or name_line['conf']<80 or
                          not same_column(name_line,job_line) or name_line['y']>job_line['y'] or len(variants)>1)
        if needs_name_retry:
            h,w=img.shape[:2]
            center=job_line['x']+job_line['w']/2
            left=min(job_line['x'],center-.20); right=max(job_line['x']+job_line['w'],center+.20)
            if name_line and name_line in name_candidates:
                left=min(left,name_line['x']); right=max(right,name_line['x']+name_line['w'])
            x1=max(0,int((left-.025)*w)); x2=min(w,int((right+.025)*w))
            y1=max(0,int((job_line['y']-.18)*h)); y2=min(h,int((job_line['y']-.004)*h))
            crop=img[y1:y2,x1:x2]
            if crop.size:
                extra=tesseract_lines(crop,7,cleaned=True)+tesseract_lines(crop,7,chroma=True)
                for ln in extra:
                    # A broken D in the common honorific Dr. is often read as O.
                    # This correction stays marked for review.
                    if re.match(r'^Or\.\s+[A-Z][a-z]+\s+[A-Z][a-z]+$',ln['text']):
                        ln['text']='Dr.'+ln['text'][3:]; ln['nameRepaired']=True
                    ln['x']=(x1+ln['x']*(x2-x1))/w; ln['y']=(y1+ln['y']*(y2-y1))/h
                    ln['w']*=((x2-x1)/w); ln['h']*=((y2-y1)/h)
                merged=merge_ocr_lines(merged+[ln for ln in extra if likely_name(ln)>0])
                name_line=max(merged,key=lambda ln:likely_name(ln)+(5 if same_column(ln,job_line) and 0<job_line['y']-ln['y']<.22 else 0))
                best=likely_name(name_line)
                name=camel_split(clean_space(name_line['text'])) if best>0 else ''
                name=re.sub(r'^[^A-Za-z]+','',name).strip()
                job,job_score,job_line=choose_job(merged)
    if retry and name_line is not None and (not job or (job_line and job_line['conf']<65)):
        # Retry below the name in its column, using the original-resolution crop.
        h,w=img.shape[:2]
        x1=max(0,int((name_line['x']-.04)*w)); x2=min(w,int((name_line['x']+name_line['w']+.04)*w))
        y1=max(0,int((name_line['y']+name_line['h'])*h)); y2=min(h,y1+int(max(.12,name_line['h']*3)*h))
        crop=img[y1:y2,x1:x2]
        if crop.size:
            extra=tesseract_lines(crop,11,cleaned=True)
            for ln in extra:
                ln['x']=(x1+ln['x']*(x2-x1))/w; ln['y']=(y1+ln['y']*(y2-y1))/h
                ln['w']*=((x2-x1)/w); ln['h']*=((y2-y1)/h)
            merged=merge_ocr_lines(merged+extra)
            job,job_score,job_line=choose_job(merged)
    raw='\n'.join(ln['text'] for ln in merged)

    # Company names may span two differently styled lines in a logo.
    domain_company=company_from_domain(website)
    domain_key=re.sub(r'[^a-z]','',domain_company.lower())
    branding=[]
    for ln in merged:
        s=clean_space(re.sub(r'^[^A-Za-z]+|[^A-Za-z.]+$','',ln['text']))
        if re.fullmatch(r'(?:[A-Za-z]\s+){3,}[A-Za-z]',s): s=s.replace(' ','')
        words=set(re.findall(r'[a-z]+',s.lower()))
        if not words or len(s)>55 or s.lower() in {name.lower(),job.lower()}: continue
        if name and re.sub(r'\W+','',name).lower() in re.sub(r'\W+','',s).lower(): continue
        if '@' in ln['text'] or any(ch.isdigit() for ch in ln['text']) or WEB_RE.search(s): continue
        if words & (LABEL_WORDS|ADDRESS_WORDS) or words & JOB_WORDS: continue
        if ln['conf']<30: continue
        branding.append({**ln,'text':s})
    candidates=list(branding)
    for a in branding:
        for b in branding:
            gap=b['y']-(a['y']+a['h'])
            if a['y']<b['y'] and -.025<=gap<=.12 and same_column(a,b):
                if len((a['text']+' '+b['text']).split())<=6:
                    candidates.append({**a,'text':a['text']+' '+b['text'],'conf':min(a['conf'],b['conf'])})
    company=''; comp_score=0; comp_line=None; company_inferred=False
    for ln in candidates:
        s=ln['text']; key=re.sub(r'[^a-z]','',s.lower()); words=set(re.findall(r'[a-z]+',s.lower()))
        similarity=SequenceMatcher(None,key,domain_key).ratio() if domain_key else 0
        hits=len(words & COMPANY_WORDS)
        if not (hits or similarity>.55 or (ln['y']<.30 and ln['h']>.025)): continue
        score=hits*5+similarity*14+min(3,ln['h']*30)+ln['conf']/35
        if score>comp_score: company=s; comp_score=score; comp_line=ln
    if domain_company:
        key=re.sub(r'[^a-z]','',company.lower())
        similarity=SequenceMatcher(None,key,domain_key).ratio()
        domain_confirmed=any(email.split('@')[-1]==website.split('://')[-1].removeprefix('www.').split('/')[0] for email in emails)
        strong_brand=comp_line and comp_line['conf']>=65 and len(company.split())>=2 and bool(set(re.findall(r'[a-z]+',company.lower()))&COMPANY_WORDS)
        if not company or similarity>.65 or (len(key)>=3 and key in domain_key) or (domain_confirmed and not strong_brand):
            company_inferred=(key!=domain_key)
            company=domain_company

    # Address: collect address-looking lines, excluding city-list marketing footer when possible
    addr=[]
    for ln in merged:
        original=clean_space(ln['text'])
        # A common OCR layout joins address (left column) with Mobile/Phone (right column). Keep the address prefix.
        s=re.split(r'\b(?:Phone|Mobile|Mob|Tel|Telephone|Email|Website)\b', original, maxsplit=1, flags=re.I)[0]
        s=clean_space(s); low=s.lower()
        if not s or s.lower() in {name.lower(),job.lower(),company.lower()}: continue
        if '@' in s or WEB_RE.search(s): continue
        words=set(re.findall(r'[a-z]+',low))
        city_hits=len(words & {'mumbai','delhi','gurgaon','gurugram','noida','bangalore','bengaluru','hyderabad','chennai','kolkata','pune','patna','raipur','jamshedpur','thane','australia','dubai'})
        has_addr=bool(words & ADDRESS_WORDS) or bool(re.search(r'\b\d{5,6}\b',s)) or (any(ch.isdigit() for ch in s) and len(s)>12)
        if not words & ADDRESS_WORDS and any(len(re.sub(r'\D','',phone))>=9 for phone in PHONE_RE.findall(s)): continue
        if has_addr and ln['y']>.25:
            if (s.count('|')>=3 or city_hits>=2) and not re.search(r'\d',s): continue
            if words & {'australia','dubai'} and not re.search(r'\d',s): continue
            # Several OCR passes can read the same address line differently.
            # Keep the strongest spatial reading instead of concatenating variants.
            score=ln['conf']+len(words & ADDRESS_WORDS)*4-2*len(re.findall(r'[^\w\s,().&/-]',s))
            duplicate=next((i for i,(old,old_ln,_) in enumerate(addr)
                            if same_column(ln,old_ln) and (
                                abs(ln['y']-old_ln['y'])<max(.025,min(ln['h'],old_ln['h'])*.6) or
                                (abs(ln['y']-old_ln['y'])<.09 and SequenceMatcher(None,s.lower(),old.lower()).ratio()>.70))),None)
            if duplicate is None: addr.append((s,ln,score))
            elif score>addr[duplicate][2]: addr[duplicate]=(s,ln,score)
    address=', '.join(dict.fromkeys(s for s,_,_ in sorted(addr,key=lambda item:item[1]['y'])))[:350]

    values={
        'name':name,'companyName':company,'jobTitle':job,'address':address,
        'email':emails[0] if emails else '','personalPhone':personal,'companyPhone':company_phone,'website':website
    }
    confidence={}
    for k,v in values.items():
        if not v: confidence[k]='missing'
        elif k=='email': confidence[k]='review' if email_repaired else 'high'
        elif k=='website': confidence[k]='high' if website_observed else 'review'
        elif k in ('personalPhone','companyPhone'): confidence[k]='high'
        elif k=='name': confidence[k]='high' if best>=12 and name_line['conf']>=65 and not name_line.get('nameRepaired') else 'review'
        elif k=='jobTitle': confidence[k]='high' if job_line and job_line['conf']>=65 and '0' not in job_line['text'] else 'review'
        elif k=='companyName': confidence[k]='high' if not company_inferred and comp_line and comp_line['conf']>=65 else 'review'
        else: confidence[k]='review'
    return values, confidence, raw


def img_to_data_url(img, max_w=650):
    h,w=img.shape[:2]
    if w>max_w:
        s=max_w/w; img=cv2.resize(img,None,fx=s,fy=s,interpolation=cv2.INTER_AREA)
    ok,b=cv2.imencode('.jpg',img,[int(cv2.IMWRITE_JPEG_QUALITY),82])
    if not ok: return ''
    return 'data:image/jpeg;base64,'+base64.b64encode(b.tobytes()).decode('ascii')


def process_image(data):
    arr=np.frombuffer(data,np.uint8); img=cv2.imdecode(arr,cv2.IMREAD_COLOR)
    if img is None: raise ValueError('Unsupported or damaged image.')
    cards=detect_cards(img)
    results=[]
    for i,crop in enumerate(cards,1):
        try:
            fields,confidence,raw=extract_fields(crop)
            results.append({'index':i,'fields':fields,'confidence':confidence,'rawText':raw,'preview':img_to_data_url(crop)})
        except Exception as e:
            results.append({'index':i,'fields':{k:'' for k in ['name','companyName','jobTitle','address','email','personalPhone','companyPhone','website']},'confidence':{},'rawText':'','preview':img_to_data_url(crop),'error':str(e)})
    return results


def export_xlsx(rows):
    wb=Workbook(); ws=wb.active; ws.title='Business Cards'
    headers=['S.No.','Name','Company Name','Job Title','Address','Email','Personal Phone','Company Phone','Website','Source Image']
    ws.append(headers)
    for cell in ws[1]:
        cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='1F4E78'); cell.alignment=Alignment(horizontal='center')
    for i,row in enumerate(rows,1):
        ws.append([i,row.get('name',''),row.get('companyName',''),row.get('jobTitle',''),row.get('address',''),row.get('email',''),row.get('personalPhone',''),row.get('companyPhone',''),row.get('website',''),row.get('sourceImage','')])
    widths=[8,24,28,28,55,34,22,22,30,28]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
    bio=io.BytesIO(); wb.save(bio); return bio.getvalue()

class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        sys.stdout.write('[server] '+fmt%args+'\n')
    def send_json(self,obj,status=200):
        b=json.dumps(obj,ensure_ascii=False).encode('utf-8'); self.send_response(status); self.send_header('Content-Type','application/json; charset=utf-8'); self.send_header('Content-Length',str(len(b))); self.end_headers()
        if self.command!='HEAD': self.wfile.write(b)
    def send_error(self, code, message=None, explain=None):
        # BaseHTTPRequestHandler normally sends HTML, including for unsupported methods.
        if urlparse(getattr(self,'path','')).path.startswith('/api/'):
            self.send_json({'ok':False,'error':message or self.responses.get(code, ('Request failed',))[0]},code)
            return
        super().send_error(code, message, explain)
    def do_GET(self):
        path=urlparse(self.path).path
        if path in ('/','/index.html'):
            p=os.path.join(ROOT,'index.html'); b=open(p,'rb').read(); self.send_response(200); self.send_header('Content-Type','text/html; charset=utf-8'); self.send_header('Content-Length',str(len(b))); self.end_headers(); self.wfile.write(b); return
        if path=='/api/health':
            self.send_json({'ok':True,'app':'business-card-digitizer-v2','ocr':'Tesseract 5 + OpenCV','multiCard':True}); return
        if path in ('/api/process','/api/export'):
            self.send_error(405,'Use POST for this API endpoint.'); return
        self.send_error(404)
    def do_POST(self):
        path=urlparse(self.path).path
        if path=='/api/health':
            self.send_error(405,'Use GET for this API endpoint.'); return
        if path not in ('/api/process','/api/export'):
            self.send_error(404); return
        try:
            n=int(self.headers.get('Content-Length','0'))
            if n>MAX_BODY:
                self.close_connection=True
                self.send_json({'ok':False,'error':'Image upload is too large. Choose a smaller photo (under 22 MB).'},413); return
            if n<=0: raise ValueError('Request body is empty.')
            body=self.rfile.read(n); data=json.loads(body.decode('utf-8'))
            if not isinstance(data,dict): raise ValueError('Expected a JSON object.')
            if path=='/api/process':
                src=data.get('data','');
                if not isinstance(src,str) or not src: raise ValueError('No image data was supplied.')
                if ',' in src: src=src.split(',',1)[1]
                raw=base64.b64decode(src,validate=True)
                if not raw: raise ValueError('No image data was supplied.')
                results=process_image(raw)
                self.send_json({'ok':True,'count':len(results),'cards':results}); return
            if path=='/api/export':
                rows=data.get('rows',[])
                if not rows: raise ValueError('No confirmed cards to export.')
                if not isinstance(rows,list) or any(not isinstance(row,dict) for row in rows): raise ValueError('Expected a list of cards to export.')
                b=export_xlsx(rows); self.send_response(200); self.send_header('Content-Type','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'); self.send_header('Content-Disposition','attachment; filename="business-cards.xlsx"'); self.send_header('Content-Length',str(len(b))); self.end_headers(); self.wfile.write(b); return
        except ValueError as e:
            self.send_json({'ok':False,'error':str(e)},400)
        except Exception as e:
            traceback.print_exc(); self.send_json({'ok':False,'error':str(e)},500)

def main():
    # helpful Windows check
    try: pytesseract.get_tesseract_version()
    except Exception:
        print('\nERROR: Tesseract OCR is not installed or not in PATH. See README.txt.\n'); raise
    host='127.0.0.1'; port=8000
    srv=ThreadingHTTPServer((host,port),Handler)
    url=f'http://{host}:{port}'
    print(f'Business Card Digitizer V2 running at {url}')
    print('Press Ctrl+C to stop.')
    threading.Timer(1.0, lambda:webbrowser.open(url)).start()
    try: srv.serve_forever()
    except KeyboardInterrupt: pass

if __name__=='__main__': main()
